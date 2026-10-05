from pathlib import Path

import pytest
from openpyxl import load_workbook

from exporter import export_to_excel
from schemas import AdvertisementDto, AdvertisementHistoryDto


def test_save_to_excel_throws_value_error_on_empty_ads_list(tmp_path: Path):
    ads = []
    with pytest.raises(ValueError):
        export_to_excel(ads, filepath=(tmp_path / "test.xlsx"))


def test_save_to_excel_throws_value_error_on_empty_path(advertisement: AdvertisementDto):
    with pytest.raises(ValueError):
        export_to_excel(ads=[advertisement], filepath=Path(""))


def test_save_to_excel_success(advertisement: AdvertisementDto, tmp_path: Path):
    export_to_excel(ads=[advertisement], filepath=tmp_path / "test.xlsx")
    assert (tmp_path / "test.xlsx").exists()


def test_export_to_excel_creates_history_sheet(
    advertisement: AdvertisementDto, adv_history: AdvertisementHistoryDto, tmp_path: Path
):
    test_filepath = tmp_path / "test.xlsx"
    export_to_excel(ads=[advertisement], ads_history=[adv_history], filepath=test_filepath)
    assert test_filepath.exists()

    wb = load_workbook(test_filepath)
    assert "Объявления" in wb.sheetnames
    assert "История изменений" in wb.sheetnames

    ws_history = wb["История изменений"]
    assert ws_history.cell(row=2, column=1).value == advertisement.kufar_id
    assert ws_history.cell(row=2, column=2).value == "title"
    assert ws_history.cell(row=2, column=3).value == "old"
    assert ws_history.cell(row=2, column=4).value == "new"
