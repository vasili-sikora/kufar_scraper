from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from config import DATA_DIR
from schemas import AdvertisementDto, AdvertisementHistoryDto

EXCEL_PATH = DATA_DIR / "advertisements.xlsx"
HEADER_FONT = Font(bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")


def _export_ads_history(wb: Workbook, ads_history: list[AdvertisementHistoryDto]) -> None:
    if not ads_history:
        return
    ws = wb.create_sheet(title="История изменений")
    headers = ["ID объявления", "Поле", "Старое значение", "Новое значение", "Дата изменения"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    for record in ads_history:
        change_date = record.changed_at.strftime("%Y-%m-%d %H:%M") if record.changed_at else ""
        ws.append(
            [
                record.kufar_id,
                record.field_name,
                record.old_value or "",
                record.new_value,
                change_date,
            ]
        )

    history_col_widths = {"A": 16, "B": 16, "C": 30, "D": 30, "E": 20}
    for col, width in history_col_widths.items():
        ws.column_dimensions[col].width = width


def _export_ads(wb: Workbook, ads: list[AdvertisementDto]) -> None:
    ws = wb.active
    if ws is None:
        raise ValueError("Something went wrong!")

    ws.title = "Объявления"
    headers = [
        "Kufar ID",
        "Название",
        "Цена",
        "Статус",
        "Описание",
        "Дата публикации",
        "Ссылка",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    for ad in ads:
        pub_date = ad.published_at.strftime("%Y-%m-%d %H:%M") if ad.published_at else ""
        ws.append(
            [
                ad.kufar_id,
                ad.title,
                ad.price,
                ad.status,
                ad.description or "",
                pub_date,
                ad.link,
            ]
        )

    column_widths = {
        "A": 14,  # Kufar ID
        "B": 35,  # Название
        "C": 12,  # Цена
        "D": 12,  # Статус
        "E": 50,  # Описание
        "F": 18,  # Дата публикации
        "G": 35,  # Ссылка
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width


def export_to_excel(
    ads: list[AdvertisementDto],
    ads_history: list[AdvertisementHistoryDto] | None = None,
    filepath: Path = EXCEL_PATH,
) -> Path:
    if not ads:
        raise ValueError("Ads list is empty!")
    if not filepath or not filepath.name:
        raise ValueError("Invalid filepath!")
    filepath.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    _export_ads(wb, ads)
    if ads_history:
        _export_ads_history(wb, ads_history)

    wb.save(filepath)
    return filepath
